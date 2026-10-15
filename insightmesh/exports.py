"""Generate PDF and Word research brief exports."""
from io import BytesIO

def pdf_export(title: str, body: str) -> bytes:
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from xml.sax.saxutils import escape
    out = BytesIO(); doc = SimpleDocTemplate(out)
    styles = getSampleStyleSheet(); story = [Paragraph(escape(title), styles["Title"]), Spacer(1, 14)]
    for line in body.splitlines():
        if line.strip(): story.append(Paragraph(escape(line[:2000]), styles["BodyText"]))
        else: story.append(Spacer(1, 8))
    doc.build(story); return out.getvalue()

def docx_export(title: str, body: str) -> bytes:
    from docx import Document
    out = BytesIO(); doc = Document(); doc.add_heading(title, 0)
    for line in body.splitlines(): doc.add_paragraph(line)
    doc.save(out); return out.getvalue()
